from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from PIL import Image
import os
from pathlib import Path

HERE = Path(__file__).resolve().parent

FIG = str(HERE / "paper_figs")
OUT = str(HERE.parent / "Gisting-CTO-deck.pptx")

BG=RGBColor(0x0e,0x14,0x17); BG2=RGBColor(0x14,0x1d,0x21); INK=RGBColor(0xee,0xf3,0xf3)
INK2=RGBColor(0x9f,0xb2,0xb6); LINE=RGBColor(0x25,0x33,0x3a); PETROL=RGBColor(0x3f,0xb0,0xc6)
PETROL2=RGBColor(0x1f,0x5f,0x6e); PLUM=RGBColor(0xc3,0x9b,0xe0); COPPER=RGBColor(0xe0,0x8a,0x4c)
GOOD=RGBColor(0x57,0xc0,0x8a); WARN=RGBColor(0xe3,0xb2,0x4c); SLATE=RGBColor(0x4a,0x5a,0x60); WHITE=RGBColor(0xff,0xff,0xff)
HEAD="Georgia"; BODY="Calibri"; MONO="Consolas"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
SW,SH=13.333,7.5

def slide():
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb=BG; r.line.fill.background()
    r.shadow.inherit=False
    return s

def tb(s,l,t,w,h,anchor=MSO_ANCHOR.TOP):
    b=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=b.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor; return tf

def setpara(p,text,size,color,bold=False,font=BODY,align=PP_ALIGN.LEFT,spacing=1.0,after=0):
    p.text=text; p.alignment=align; p.line_spacing=spacing; p.space_after=Pt(after)
    for r in p.runs:
        r.font.size=Pt(size); r.font.bold=bold; r.font.name=font; r.font.color.rgb=color
    return p

def addpara(tf,text,size,color,bold=False,font=BODY,align=PP_ALIGN.LEFT,spacing=1.0,after=0,bullet=False):
    p=tf.add_paragraph(); return setpara(p,text,size,color,bold,font,align,spacing,after)

def first(tf,text,size,color,bold=False,font=BODY,align=PP_ALIGN.LEFT,spacing=1.0,after=0):
    return setpara(tf.paragraphs[0],text,size,color,bold,font,align,spacing,after)

def rrect(s,l,t,w,h,fill=BG2,line=LINE,lw=1.0,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh=s.shapes.add_shape(shape,Inches(l),Inches(t),Inches(w),Inches(h))
    sh.shadow.inherit=False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    try: sh.adjustments[0]=0.06
    except Exception: pass
    return sh

def eyebrow(s,text,y=0.5):
    first(tb(s,0.7,y,11.9,0.4),text.upper(),13,PETROL,bold=True,font=MONO)
def title(s,text,y=0.92,size=34,w=11.9):
    first(tb(s,0.7,y,w,1.5),text,size,INK,bold=True,font=HEAD,spacing=1.02)
def pic(s,path,l,t,w):
    im=Image.open(path); ar=im.height/im.width; ph=w*ar
    # white card behind
    rrect(s,l-0.12,t-0.12,w+0.24,ph+0.24,fill=WHITE,line=LINE,lw=1.0)
    s.shapes.add_picture(path,Inches(l),Inches(t),Inches(w))
    return ph

def card(s,l,t,w,h,kicker,kcolor,bodyfn):
    rrect(s,l,t,w,h,fill=BG2,line=LINE,lw=1.0)
    tf=tb(s,l+0.28,t+0.24,w-0.56,h-0.48)
    first(tf,kicker.upper(),11,kcolor,bold=True,font=MONO,after=6)
    bodyfn(tf)

PLUMLN=RGBColor(0x5a,0x3f,0x6e); BADLN=RGBColor(0x6e,0x3a,0x3a); OKLN=RGBColor(0x2c,0x50,0x40); WARNLN=RGBColor(0x5a,0x4a,0x24)
def cards(s,items,y,h,n,size=15,cw=None,gap=0.2):
    cw=cw or (11.9-gap*(n-1))/n
    for i,it in enumerate(items):
        k,c,body=it[:3]; ln=it[3] if len(it)>3 else LINE
        l=0.7+i*(cw+gap); rrect(s,l,y,cw,h,fill=BG2,line=ln,lw=1.25 if ln!=LINE else 1.0)
        tf=tb(s,l+0.28,y+0.24,cw-0.56,h-0.48)
        first(tf,k.upper(),11,c,bold=True,font=MONO,after=6)
        if len(it)>4: addpara(tf,it[4],11,INK2,font=MONO,after=6)
        addpara(tf,body,size,INK,spacing=1.15)
def caption(s,text,y=6.3,size=13):
    first(tb(s,0.7,y,11.9,0.8),text,size,INK2,spacing=1.15)
def tag(s,text,y,w=4.6,col=WARN,ln=WARNLN):
    rrect(s,0.7,y,w,0.5,fill=None,line=ln,lw=1.0); first(tb(s,0.85,y+0.06,w-0.2,0.4),text,12,col,font=MONO)

# 1 TITLE
s=slide()
eyebrow(s,"Lattice · cost optimization · new pillar",0.9)
first(tb(s,0.7,1.35,11.5,2.6),"Steno: learned prompt compression for agent calls",44,INK,bold=True,font=HEAD,spacing=1.02)
first(tb(s,0.7,4.2,10.5,1.3),"A learned shorthand for the fixed preamble a coding agent re-sends to the model: the technique Shopify calls gisting. Shown on one harness and model pair; the next step tests whether it carries to others.",18,INK2,spacing=1.2)
tf=tb(s,0.7,6.0,11.5,0.9)
first(tf,"PAI",13,INK2,font=MONO)

# 2 BLUF
s=slide(); eyebrow(s,"Bottom line up front"); title(s,"The answer in three lines")
cards(s,[("Where it fits",PETROL,"A new pillar under Lattice, alongside the meta-harness, adaptive routing and model distillation tracks. Steno shortens the fixed prompt a coding agent re-sends to the model."),
         ("What we showed",PLUM,"On Claude Code with Qwen at 8:1, input per turn fell from 24.3k to 9.4k tokens with scores matching the full prompt on our suites. Short H100 replays: 2x the request rate within each arm’s latency threshold, +44% peak throughput."),
         ("The ask",GOOD,"Run the same experiment suite on the next harness and distilled-model pair, validate on held-out work, then a guarded pilot.")],2.1,4.3,3,size=15)

# 3 LATTICE
s=slide(); eyebrow(s,"Lattice · cost optimization"); title(s,"A new pillar under Lattice")
first(tb(s,0.7,2.0,11.6,1.0),"Steno targets serving cost by shortening repeated prompts. Short H100 replays showed higher throughput; savings per successful task are not yet proven.",17,INK2,spacing=1.2)
cw=(11.9-0.6)/4
for i,(k,c,ln,body) in enumerate([("Meta-harness",PETROL,LINE,""),("Adaptive routing",PETROL,LINE,""),("Model distillation",PETROL,LINE,""),("Steno · new",PLUM,PLUMLN,"This deck.")]):
    l=0.7+i*(cw+0.2); rrect(s,l,3.3,cw,1.3,fill=BG2,line=ln,lw=1.25 if ln!=LINE else 1.0)
    tf=tb(s,l+0.28,3.5,cw-0.56,1.0); first(tf,k.upper(),12,c,bold=True,font=MONO,after=6)
    if body: addpara(tf,body,15,INK)

# 3b WHAT STENO IS
s=slide(); eyebrow(s,"What Steno is made of"); title(s,"Five parts, designed for any harness and model pair")
cards(s,[("1 · Span analysis",COPPER,"A harness adapter per harness feeds one shared analysis. It checks every call, splits the fixed part from per-session values, and fails on anything it cannot explain.",RGBColor(0x5a,0x40,0x2a)),
         ("2 · Trainer",PLUM,"Adds new token rows to the model and trains them by self-distillation, with the base model frozen.",PLUMLN),
         ("3 · Proxy",PETROL,"Swaps the fixed span for the Steno tokens and adopts the same harness adapter. No change to the agent or the engine’s code; the served model carries the new token rows.")],2.0,2.1,3,size=14)
cards(s,[("4 · Benchmark pack",PETROL,"Task suites scored against the full prompt, plus a serving benchmark for throughput and latency."),
         ("5 · Auto loop",GOOD,"Runs the other four as stages. Runs today: one run spec, span study first, budget and pass gates, verified teardown. Bounded self-improvement (RSI, time-boxed): failure triage, lessons turned into checks, gated promotion, and a next-recipe proposer a person approves.",OKLN)],4.35,1.95,2,size=14)

# 4 SPAN + RECIPE
s=slide(); eyebrow(s,"The span, and why it is per pair"); title(s,"What the agent re-sends, measured on one pair")
by,bh,total=1.95,0.95,11.9; cx=0.7
for name,sub,frac,col in [("Tool schemas","~16,000 tokens  ·  >90% of the block",0.78,PETROL2),("Rules","~1.3k",0.09,COPPER),("Per-session","kept raw",0.13,RGBColor(0x2b,0x3a,0x41))]:
    w=total*frac; rrect(s,cx,by,w-0.03,bh,fill=col,line=None,shape=MSO_SHAPE.RECTANGLE)
    tf=tb(s,cx+0.18,by,w-0.36,bh,anchor=MSO_ANCHOR.MIDDLE)
    first(tf,name,14 if frac>0.2 else 11,WHITE,bold=True,after=2); addpara(tf,sub,10 if frac>0.2 else 8.5,WHITE)
    cx+=w
for i,(n,c,u) in enumerate([("17.5–21k",COPPER,"fixed tokens per call, across the span configurations we measured"),("0.72",PETROL,"mean share of input across eight logged sessions")]):
    first(tb(s,0.7+i*5.5,3.05,5.2,0.6),n,26,c,bold=True,font=MONO); first(tb(s,0.7+i*5.5,3.62,5.2,0.5),u,11,INK2,font=MONO)
cards(s,[("Rerun per pair",COPPER,"The whole experiment suite, span study first, runs again for each harness and distilled-model pair. Needs a model whose weights we host.",RGBColor(0x5a,0x40,0x2a))],4.3,1.4,1,size=15)
caption(s,"Measured on Claude Code with the Qwen tokenizer only. For each new pair, measure the repeated prompt and check which session-specific values must remain raw.",6.15,12.5)

# 4b ONBOARD A HARNESS
s=slide(); eyebrow(s,"Bringing a new harness onto Steno"); title(s,"One adapter per harness, including our own")
y,h=2.0,1.15; cx=0.7
for i,(n,w,ln,tc,sub) in enumerate([("Harness traffic",2.2,LINE,WHITE,None),("Harness adapter",2.5,RGBColor(0x5a,0x40,0x2a),COPPER,"per harness"),("Shared span analysis",2.9,LINE,WHITE,"same for all"),("Model adapter",2.4,PLUMLN,PLUM,"per model")]):
    rrect(s,cx,y,w,h,fill=BG2,line=ln); tf=tb(s,cx,y,w,h,anchor=MSO_ANCHOR.MIDDLE); first(tf,n,14,tc,bold=True,align=PP_ALIGN.CENTER)
    if sub: addpara(tf,sub,10,INK2,font=MONO,align=PP_ALIGN.CENTER)
    cx+=w
    if i<3: first(tb(s,cx,y,0.45,h,anchor=MSO_ANCHOR.MIDDLE),"→",20,PETROL,bold=True,align=PP_ALIGN.CENTER); cx+=0.45
cards(s,[("1 · Capture",PETROL,"Record a few real sessions of the harness."),
         ("2 · Adapter",PETROL,"Map its requests to the common call record, and declare its per-session values."),
         ("3 · Fixtures",PETROL,"Commit sample requests with their expected parse."),
         ("4 · Checks",GOOD,"Pass the every-call invariance check on a fresh capture, then token parity once the model adapter is in.",OKLN)],3.5,2.2,4,size=13.5)
first(tb(s,0.7,5.85,11.9,0.5),"Already earning its keep: on real Claude Code traffic, the every-call check caught 15 per-call differences the old rules missed.",15,INK,spacing=1.15)
caption(s,"Everything else in Steno stays unchanged. This is how in-house PayPal harnesses come onto Steno with little effort.",6.45,12)

# 5 PROOF
s=slide(); eyebrow(s,"Results on the pair studied · Claude Code × Qwen3.8"); title(s,"The short prompt did the same work")
for i,(r,sc) in enumerate([("2:1","12 / 12"),("4:1","12 / 12"),("8:1","12 / 12"),("16:1","12 / 12")]):
    rrect(s,0.7+i*2.2,2.0,1.9,1.1,fill=BG2,line=PETROL2,lw=1.25); tf=tb(s,0.7+i*2.2,2.05,1.9,1.0,anchor=MSO_ANCHOR.MIDDLE)
    first(tf,r,22,PETROL,bold=True,font=MONO,align=PP_ALIGN.CENTER,after=2); addpara(tf,sc,13,GOOD,font=MONO,align=PP_ALIGN.CENTER)
first(tb(s,0.7,3.3,11.6,1.3),"All four ratios scored 12/12 on the easy suite, matching the full prompt. At 8:1, input per turn fell from 24.3k to 9.4k tokens. We selected 8:1: the tested 16:1 checkpoint delivered about 61% of its peak-throughput improvement.",16,INK2,spacing=1.2)
first(tb(s,0.7,4.6,11.6,1.3),"On the harder suite, keeping session values such as the working path out of the compressed prompt, together with retraining, brought 8:1 to 16/16, matching the full prompt on the same host. Other run conditions also changed, and one path error remained.",16,INK2,spacing=1.2)
tag(s,"single runs · small, partly reused suites",6.1,4.8)

# 6 PAYOFF
s=slide(); eyebrow(s,"The payoff"); title(s,"More work from the same GPU")
cd=CategoryChartData(); cd.categories=["1","2","4","8","16","32","48","64"]
cd.add_series("full prompt",(13.1,22.7,36.9,34.7,42.7,42.7,32.0,29.3)); cd.add_series("Steno 8:1",(14.7,25.3,42.7,48.0,59.3,61.3,52.7,41.3))
gf=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(0.7),Inches(2.1),Inches(7.4),Inches(4.6),cd)
ch=gf.chart; ch.has_title=True; ch.chart_title.text_frame.text="requests / minute vs concurrent sessions  ·  mean of 3 repeats"
for r in ch.chart_title.text_frame.paragraphs[0].runs: r.font.size=Pt(12); r.font.color.rgb=INK2; r.font.name=MONO; r.font.bold=False
ch.has_legend=True; ch.legend.position=XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout=False; ch.legend.font.color.rgb=INK2; ch.legend.font.size=Pt(12)
ser=ch.plots[0].series
ser[0].format.line.color.rgb=SLATE; ser[0].format.line.width=Pt(3); ser[1].format.line.color.rgb=PLUM; ser[1].format.line.width=Pt(3)
for ax in (ch.category_axis,ch.value_axis):
    ax.tick_labels.font.color.rgb=INK2; ax.tick_labels.font.size=Pt(12); ax.format.line.color.rgb=LINE
ch.value_axis.has_major_gridlines=True; ch.value_axis.major_gridlines.format.line.color.rgb=RGBColor(0x1b,0x26,0x2b)
tf=tb(s,8.4,2.1,4.3,4.9)
first(tf,"2x",28,PLUM,bold=True,font=MONO,after=2); addpara(tf,"highest tested arrival rate passing each arm’s own latency threshold (18 → 36 req/min)",11.5,INK2,font=MONO,after=10)
addpara(tf,"+44%",28,PETROL,bold=True,font=MONO,after=2); addpara(tf,"peak throughput, same H100 (42.7 → 61.3 req/min)",11.5,INK2,font=MONO,after=10)
addpara(tf,"A capacity lever, not a speed lever: a single reply is barely faster.",14,INK,spacing=1.15,after=8)
addpara(tf,"Threshold: twice each arm’s unloaded median latency, about 8.5 s full and 8.0 s Steno. Short, fixed-output replays on one GPU class, mean of three repeats; they did not measure task quality at load or isolate prefix-cache effects. On a larger H200, Steno was level at peak and ahead under heavy load (next slide).",10.5,INK2,spacing=1.12)

# 6b H200
s=slide(); eyebrow(s,"Second GPU · H200 NVL, 143 GB"); title(s,"On a larger GPU, the gain shows up under heavy load")
cw=(11.9-0.4)/3
for i,(hdr,ln,vals) in enumerate([("At each arm’s peak",LINE,[("85.3",PETROL,"full prompt"),("83.3",PLUM,"Steno 8:1 · level")]),
                                  ("128 sessions · cap lifted",PLUMLN,[("47.0",SLATE,"full prompt"),("80.7",PLUM,"Steno 8:1 · 1.7x")]),
                                  ("256 sessions · cap lifted",PLUMLN,[("46.3",SLATE,"full prompt"),("56.7",PLUM,"Steno 8:1 · 1.2x")])]):
    l=0.7+i*(cw+0.2); rrect(s,l,2.0,cw,1.9,fill=BG2,line=ln)
    first(tb(s,l+0.25,2.15,cw-0.5,0.4),hdr.upper(),10.5,PETROL,bold=True,font=MONO)
    for j,(n,c,u) in enumerate(vals):
        first(tb(s,l+0.25+j*1.85,2.6,1.8,0.7),n,28,c,bold=True,font=MONO); first(tb(s,l+0.25+j*1.85,3.3,1.8,0.4),u,10,INK2,font=MONO)
first(tb(s,0.7,4.0,11.9,0.4),"replay requests per minute",11,INK2,font=MONO)
first(tb(s,0.7,4.45,11.9,1.4),"At each arm’s best point the two run level. Pushed to 128 and 256 concurrent sessions with our client’s connection cap removed, the full prompt fell to about 47 req/min as the cache filled, while Steno served 1.2 to 1.7 times more. Per rental dollar, the H100 with Steno and the H200 without it come out roughly even (appendix B).",15,INK2,spacing=1.2)
tag(s,"peak: two repeats, H100-tuned settings · heavy load: one run per condition",6.05,7.4)

# 7 SYNERGY
s=slide(); eyebrow(s,"Synergy within Lattice · proposal"); title(s,"Distil first, then apply Steno to it")
y,h=2.3,1.15; cx=0.7
nodes=[("Base model",1.8,LINE,WHITE,None),("Distil",1.6,LINE,WHITE,"distillation track"),("Student model",2.0,LINE,WHITE,None),("Steno",1.5,PLUMLN,PLUM,"this recipe"),("Smaller model, shorter prompt",2.6,LINE,WHITE,None)]
for i,(n,w,ln,tc,sub) in enumerate(nodes):
    rrect(s,cx,y,w,h,fill=BG2,line=ln); tf=tb(s,cx,y,w,h,anchor=MSO_ANCHOR.MIDDLE); first(tf,n,14,tc,bold=True,align=PP_ALIGN.CENTER)
    if sub: addpara(tf,sub,10,INK2,font=MONO,align=PP_ALIGN.CENTER)
    cx+=w
    if i<len(nodes)-1: first(tb(s,cx,y,0.45,h,anchor=MSO_ANCHOR.MIDDLE),"→",20,PETROL,bold=True,align=PP_ALIGN.CENTER); cx+=0.45
first(tb(s,0.7,3.9,11.6,1.3),"A distilled student model still receives the fixed preamble on each call. Running the Steno recipe on it aims to reduce both model size and prompt length.",18,INK2,spacing=1.25)
tag(s,"proposal · not yet tested",5.5,3.4)

# 8 JETSTREAM
s=slide(); eyebrow(s,"Where it could land · proposal"); title(s,"A first home in the Jetstream inner loop")
first(tb(s,0.7,2.2,11.2,1.5),"The Jetstream inner loop takes a ticket to a PR through coding agents on different harnesses. We propose it as the first place to trial Steno, one harness and model pair at a time.",19,INK2,spacing=1.25)
tag(s,"proposal",4.0,1.6)
caption(s,"Steno applies where the agent calls a model whose weights we host.",4.8)

# 9 LEDGER
s=slide(); eyebrow(s,"Honest ledger"); title(s,"What’s proven, what isn’t")
def col(l,kicker,kc,items,linec):
    rrect(s,l,2.15,5.75,4.5,fill=BG2,line=linec,lw=1.0)
    tf=tb(s,l+0.3,2.4,5.15,4.1); first(tf,kicker.upper(),12,kc,bold=True,font=MONO,after=10)
    for it in items: addpara(tf,"▪  "+it,13.5,INK,spacing=1.1,after=7)
col(0.7,"Established on this pair",GOOD,["Easy suite: 12/12 at every ratio. Hard suite: corrected 8:1 matched the full prompt, 16/16.","At 8:1, input per turn down from 24.3k to 9.4k tokens.","Short H100 replays: 2x rate within each arm’s latency threshold, +44% peak.","A spend-safe lab that ran the recipe end to end."],OKLN)
col(6.85,"Not yet",WARN,["The experiment suite on a second harness and distilled-model pair.","Quality on held-out work, and at production load.","Savings per successful task, and why the gain happens. An external review withdrew our first explanation.","Audit that compressed rules still bind."],WARNLN)

# 10 ASK
s=slide(); eyebrow(s,"The ask"); title(s,"Make Steno a Lattice pillar")
cards(s,[("1 · Next pair",PLUM,"Write the harness adapter, then run the experiment suite, span study first, on the next harness and distilled-model pair. Tests whether the recipe carries.",PLUMLN),
         ("2 · Validate",PETROL,"Held-out tasks, repeated runs, and quality checked at production load."),
         ("3 · Guarded pilot",PETROL,"Subject to validation, a guarded trial in the Jetstream inner loop.")],2.1,2.5,3,size=15)
first(tb(s,0.7,4.9,11.0,1.1),"The lever is real on one pair. Next, test whether it carries to a distilled model on another harness.",24,WHITE,bold=True,font=HEAD,spacing=1.1)
first(tb(s,0.7,6.35,11.9,0.5),"White paper, code and results: GitHub arunmenon/gisting-coding-agent  ·  weights and data on Hugging Face (private)",12,PETROL,font=MONO)

# APPENDIX A
s=slide(); eyebrow(s,"Appendix A · a correction we made"); title(s,"Our first explanation of the gain was wrong")
first(tb(s,0.7,2.0,11.6,1.5),"External review found a 100-connection limit in our own load generator; that limit, not the GPU, produced the number our explanation rested on. A follow-up confirmed it and we withdrew the explanation. The peak difference, measured well below that limit, survives the correction; the other measurement limits still apply.",15.5,INK2,spacing=1.2)
gx,gy,cell,gp=0.7,3.75,0.42,0.1
for k in range(64):
    rrect(s,gx+(k%16)*(cell+gp),gy+(k//16)*(cell+gp),cell,cell,fill=PETROL if k<16 else RGBColor(0x24,0x30,0x36),line=None)
caption(s,"The model has 16 full-attention and 48 linear-attention layers; their contribution to the observed throughput difference has not been isolated.",6.0)

# APPENDIX B
s=slide(); eyebrow(s,"Appendix B · cost per rental dollar"); title(s,"Roughly even across the two cards")
for i,(n,c,u) in enumerate([("16.16",SLATE,"H100 · full prompt"),("23.23",PLUM,"H100 · Steno 8:1"),("23.38",PETROL,"H200 · full prompt"),("22.83",SLATE,"H200 · Steno 8:1")]):
    first(tb(s,0.7+i*3.0,2.2,2.8,0.8),n,34,c,bold=True,font=MONO); first(tb(s,0.7+i*3.0,2.95,2.8,0.5),u,12,INK2,font=MONO)
first(tb(s,0.7,3.55,11.9,0.5),"peak replay requests per minute per dollar-hour, quoted prices: H100 $2.64/h, H200 $3.65/h",12,INK2,font=MONO)
first(tb(s,0.7,4.3,11.9,1.3),"On the H100 Steno lifts peak replay output per dollar by about 44%. The H100 with Steno lands level with the H200 without it, so this does not show that a cheaper card can replace a pricier one.",17,INK,spacing=1.25)

# APPENDIX C
s=slide(); eyebrow(s,"Appendix C · the failure we caught"); title(s,"What to keep out of the compressed prompt")
def pill(l,t,w,txt,tc,lc,h=0.82):
    rrect(s,l,t,w,h,fill=BG2,line=lc); first(tb(s,l+0.2,t,w-0.4,h,anchor=MSO_ANCHOR.MIDDLE),txt,13,tc,font=MONO,spacing=1.05)
def arrow(l,t,a="→",c=PETROL,w=0.5,h=0.82):
    first(tb(s,l,t,w,h,anchor=MSO_ANCHOR.MIDDLE),a,20,c,bold=True,align=PP_ALIGN.CENTER)
first(tb(s,0.7,2.55,1.4,0.5),"BEFORE",13,WARN,bold=True,font=MONO)
pill(2.2,2.4,4.3,"Steno: rules + …/<session-id>/",PLUM,PLUMLN); arrow(6.6,2.4)
pill(7.2,2.4,5.1,"writes to an invented directory   ×",RGBColor(0xe7,0x91,0x91),BADLN)
first(tb(s,0.7,3.75,1.4,0.5),"AFTER",13,GOOD,bold=True,font=MONO)
pill(2.2,3.6,2.3,"Steno: rules",PLUM,PLUMLN); arrow(4.55,3.6,"+",INK2,0.4)
pill(5.0,3.6,3.4,"raw: session path, date, model",INK,LINE); arrow(8.5,3.6)
pill(9.1,3.6,3.2,"writes to the intended path   ✓",GOOD,OKLN)
first(tb(s,0.7,4.85,11.9,1.1),"After keeping session values raw and retraining, corrected 8:1 scored 16/16, up from 11/16 (15/15 from 11/15 excluding one defective probe). Other run conditions also changed, and one invented-path event remained.",16,INK,spacing=1.2)
caption(s,"For each new pair, check which session-specific values must remain raw.",6.2)

# APPENDIX D
s=slide(); eyebrow(s,"Appendix D · the lab"); title(s,"The lab that ran this study")
y,h=2.6,1.05; cx=0.7
nodes=[("Provision GPU",2.0),("Train",1.5),("Serve",1.5),("Evaluate",1.7),("Destroy",1.6)]
for i,(n,w) in enumerate(nodes):
    rrect(s,cx,y,w,h,fill=BG2,line=LINE); first(tb(s,cx,y,w,h,anchor=MSO_ANCHOR.MIDDLE),n,14,WHITE,bold=True,align=PP_ALIGN.CENTER); cx+=w
    if i<len(nodes)-1: first(tb(s,cx,y,0.55,h,anchor=MSO_ANCHOR.MIDDLE),"→",20,PETROL,bold=True,align=PP_ALIGN.CENTER); cx+=0.55
rrect(s,0.7,4.1,11.9,0.95,fill=RGBColor(0x22,0x1a,0x11),line=RGBColor(0x5a,0x4a,0x2a))
first(tb(s,0.95,4.1,11.4,0.95,anchor=MSO_ANCHOR.MIDDLE),"SPEND-SAFETY · structural  ·  idle watchdog · ledger at creation · sync-before-destroy · global deadline",13,COPPER,font=MONO)
caption(s,"Most experiments after the first ran through this controller rather than by hand. The loop reaches compute through a provider-neutral interface; vast.ai is the first backend.",5.4,14)


prs.save(OUT)
print("saved",OUT,"slides",len(prs.slides._sldIdLst))
