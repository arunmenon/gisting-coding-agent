import base64, io, re
from pathlib import Path

HERE = Path(__file__).resolve().parent
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from PIL import Image

soup = BeautifulSoup(open(HERE / "gisting-neurips.html").read(), "html.parser")
wrap = soup.select_one(".wrap")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Georgia"; st.font.size = Pt(10.5)
INK2 = RGBColor(0x56,0x5b,0x62); PETROL = RGBColor(0x1f,0x5f,0x6e)

def add_inline(par, node):
    for c in node.children:
        if isinstance(c, NavigableString):
            par.add_run(str(c))
        elif isinstance(c, Tag):
            if c.name in ("b","strong"):
                r = par.add_run(c.get_text()); r.bold = True
            elif c.name in ("em","i"):
                r = par.add_run(c.get_text()); r.italic = True
            elif c.name == "code":
                r = par.add_run(c.get_text()); r.font.name = "Consolas"; r.font.size = Pt(9)
            elif c.name == "span":  # section number
                r = par.add_run(c.get_text()+" "); r.font.color.rgb = PETROL; r.font.name="Consolas"; r.bold=True
            else:
                add_inline(par, c)

def emit(node):
    cls = node.get("class", []) if isinstance(node, Tag) else []
    name = node.name
    if name == "h1":
        p = doc.add_heading("", level=0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(node.get_text()); r.font.size = Pt(17)
    elif name == "div" and "byline" in cls:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(node.get_text()); r.font.color.rgb = INK2; r.font.size = Pt(10)
    elif name == "div" and "venue" in cls:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(node.get_text()); r.font.color.rgb = INK2; r.font.size = Pt(8)
    elif name == "div" and "abs" in cls:
        for ch in node.find_all("p", recursive=False):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.right_indent = Inches(0.3)
            add_inline(p, ch)
    elif name == "h2":
        p = doc.add_heading(level=1); add_inline(p, node)
    elif name == "h3":
        p = doc.add_heading(level=2); add_inline(p, node)
    elif name == "p":
        p = doc.add_paragraph()
        if "tcap" in cls or "contrib" in cls:
            add_inline(p, node)
            for r in p.runs:
                if "tcap" in cls: r.font.size = Pt(9); r.font.color.rgb = INK2
        else:
            add_inline(p, node)
    elif name in ("ul","ol"):
        for index, li in enumerate(node.find_all("li", recursive=False), start=1):
            p = doc.add_paragraph(style="List Bullet" if name=="ul" else "Normal")
            if name == "ol":
                # Keep each HTML list independent, including references after contributions.
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.add_run(f"{index}. ")
            add_inline(p, li)
    elif name == "figure":
        img = node.find("img")
        if img and img.get("src","").startswith("data:image"):
            b64 = img["src"].split(",",1)[1]
            raw = base64.b64decode(b64)
            im = Image.open(io.BytesIO(raw)); w,h = im.size
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(io.BytesIO(raw), width=min(Inches(6.3), doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin))
            p.paragraph_format.keep_with_next = True
        cap = node.find("figcaption")
        if cap:
            p = doc.add_paragraph(); add_inline(p, cap)
            for r in p.runs: r.font.size = Pt(9); 
    elif name == "div" and "tw" in cls:
        emit_table(node.find("table"))
    elif name == "table":
        emit_table(node)
    elif name == "footer":
        doc.add_paragraph()
        for ch in node.children:
            if isinstance(ch, Tag):
                emit(ch)
            elif isinstance(ch, NavigableString) and str(ch).strip():
                doc.add_paragraph(str(ch).strip())
    elif name == "div":
        for ch in node.children:
            if isinstance(ch, Tag): emit(ch)

def emit_table(tbl):
    if tbl is None: return
    rows = tbl.find_all("tr")
    ncol = max(len(r.find_all(["td","th"])) for r in rows)
    t = doc.add_table(rows=0, cols=ncol); t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, r in enumerate(rows):
        cells = r.find_all(["td","th"])
        table_row = t.add_row()
        row = table_row.cells
        properties = table_row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index == 0 and all(c.name == "th" for c in cells):
            properties.append(OxmlElement("w:tblHeader"))
        for i,c in enumerate(cells):
            row[i].text = ""
            p = row[i].paragraphs[0]; rn = p.add_run(c.get_text())
            rn.font.size = Pt(8.5)
            if c.name == "th": rn.bold = True

for node in wrap.children:
    if isinstance(node, Tag):
        emit(node)

out = HERE.parent / "Gisting-NeurIPS-paper.docx"
doc.save(out)
print("saved", out, "| images", len(doc.inline_shapes), "| tables", len(doc.tables))
assert len(doc.inline_shapes) == len(wrap.select("figure img"))
assert len(doc.tables) == len(wrap.select("table"))
